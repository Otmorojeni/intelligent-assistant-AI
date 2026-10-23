"""Unit and integration test suite for DOCX Document Generator (doc_generator.py)

Covers Section 3.3 document cleaning, Section 2.3 Project Card creation,
note_text formatting, sequential approval routing, and FastAPI REST endpoints.
"""

import asyncio
from datetime import datetime
from pathlib import Path
import re
import sys
import unittest
import uuid

# Ensure project root is in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

# Ensure UTF-8 output
if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import docx
import httpx
from httpx import ASGITransport
from sqlalchemy import select

from src.config import settings

from src.core.doc_generator import (
    clean_and_populate_docx,
    clean_text_strict,
    find_template_file,
    generate_document,
    get_all_projects,
    get_available_templates,
    get_template_by_code,
)
from src.db.database import async_session_factory
from src.db.models import AuditLog, ProjectCard, TemplateModule
from src.main import app
from src.schemas.api_models import DocumentGenerateRequest

SAMPLE_MODULE_FIELDS = {
    "MOD-ORDER-IS": {
        "system_name": "АСУ ДОУ «Приоритет»",
        "purpose": "автоматизации сквозного учета и согласования документов",
        "commissioning_date": "01.11.2026",
        "responsible_dept": "Департамент информационных технологий",
        "infosec_dept": "Департамент информационной безопасности",
        "control_officer": "директора Департамента информационных технологий",
    },
    "MOD-ORDER-DOU": {
        "department_name": "Департамент корпоративного управления",
        "primary_responsible": "главный эксперт Соколова М.В.",
        "deputy_responsible": "ведущий советник Куликов А.П.",
        "reason": "Инструкции по делопроизводству № ОД-184",
        "control_officer": "руководителя Службы ДОУ",
    },
    "MOD-ORDER-PROP": {
        "department_name": "Управление безопасности",
        "transferor_name": "начальник управления Волков В.В.",
        "transferee_name": "заместитель начальника Морозов М.М.",
        "transfer_basis": "плановой ротацией кадров",
        "deadline": "15.11.2026",
        "commission_members": "начальник службы кадров",
        "affects_is_access": True,
        "control_officer": "руководителя аппарата",
    },
    "MOD-ORDER-LNA": {
        "act_title": "Регламент защиты конфиденциальной информации",
        "purpose": "повышения уровня информационной безопасности",
        "effective_date": "10.11.2026",
        "responsible_dept": "Департамент информационной безопасности",
        "placement_location": "на внутреннем корпоративном портале",
        "repeal_act": "",
        "control_officer": "директора Департамента ИБ",
    },
    "MOD-DISP-AUDIT": {
        "audit_process": "управления операционными рисками",
        "audit_goal": "оценки соответствия локальным нормативным актам",
        "audit_period": "с 05.11.2026 по 25.11.2026",
        "audited_dept": "Служба управления рисками",
        "audit_leader": "главный аудитор Новиков Н.Н.",
        "audit_team": "старший аудитор Лебедев Л.Л.",
        "report_deadline": "30.11.2026",
        "control_officer": "директора Департамента внутреннего аудита",
    },
    "MOD-DISP-PROJECT": {
        "project_name": "Единая цифровая платформа Банка",
        "committee_purpose": "координации внедрения цифровых сервисов",
        "reason": "стратегией цифровизации на 2026-2028 годы",
        "managing_dept": "Проектный офис",
        "chairman": "Заместитель Председателя Организации",
        "deputy_chairman": "Директор Департамента ИТ",
        "secretary": "руководитель проекта ИТ",
        "members": "руководители департаментов-заказчиков",
        "is_it_project": True,
        "repeal_order": "",
        "control_officer": "Председателя Организации",
    },
    "MOD-DISP-WORKGROUP": {
        "workgroup_subject": "подготовке новой редакции учетной политики",
        "goal": "актуализации правил бухгалтерского и налогового учета",
        "leader": "главный бухгалтер Ковалева Е.И.",
        "deputy_leader": "заместитель главного бухгалтера",
        "secretary": "ведущий бухгалтер Смирнова О.С.",
        "members": "бухгалтеры профильных отделов",
        "deadline": "20.12.2026",
        "expected_result": "проект учетной политики на 2027 год",
        "termination_event": "утверждения учетной политики",
        "is_it_project": False,
        "control_officer": "Заместителя Председателя Организации",
    },
}


class TestDocGenerator(unittest.IsolatedAsyncioTestCase):
    """Test suite for DOCX Document Generator, cleaning, DB persistence, and REST endpoints."""

    def test_clean_text_strict_unit(self):
        """Verify strict regex cleaning: footnotes, lines, empty brackets, and service zones."""
        # 1. Footnote removal
        raw = "Утвердить положение [2] согласно приложению."
        cleaned = clean_text_strict(raw)
        self.assertNotIn("[2]", cleaned)
        self.assertEqual(cleaned, "Утвердить положение согласно приложению.")

        # 2. Underscores removal
        raw_lines = "О вводе в эксплуатацию ____________________ [1]"
        cleaned_lines = clean_text_strict(raw_lines)
        self.assertNotIn("___", cleaned_lines)
        self.assertNotIn("[1]", cleaned_lines)

        # 3. Preservation of service zones
        sz1 = "ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)"
        sz2 = "ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)"
        self.assertEqual(clean_text_strict(sz1), sz1)
        self.assertEqual(clean_text_strict(sz2), sz2)

        # 4. Signatory replacement
        sig_raw = "И. О. Фамилия [15]"
        sig_cleaned = clean_text_strict(sig_raw, signatory_name="А.В. Смирнов")
        self.assertIn("А.В. Смирнов", sig_cleaned)
        self.assertNotIn("[15]", sig_cleaned)

    async def test_generate_all_7_templates_strict_cleaning(self):
        """Generate all 7 modular templates and verify strict cleaning per Section 3.3."""
        async with async_session_factory() as session:
            templates = await get_available_templates(session)
            self.assertEqual(len(templates), 7, "Registry must contain exactly 7 templates")

            for tm in templates:
                with self.subTest(module_code=tm.module_code):
                    fields = SAMPLE_MODULE_FIELDS.get(tm.module_code, {})
                    req = DocumentGenerateRequest(
                        module_code=tm.module_code,
                        executor_name="Иванов И.И.",
                        executor_department="Департамент информационных технологий",
                        signatory_name="Петров П.П.",
                        access_level="public",
                        fields=fields,
                        user_role="employee",
                    )
                    resp = await generate_document(session, req)
                    self.assertTrue(resp.file_name.endswith(".docx"))

                    out_path = settings.GENERATED_DOCS_DIR / resp.file_name
                    self.assertTrue(out_path.exists(), f"File {out_path} must exist on disk")

                    # Open generated DOCX and extract all text
                    doc = docx.Document(str(out_path))
                    para_texts = [p.text for p in doc.paragraphs if p.text.strip()]
                    tbl_texts = [
                        c.text
                        for tbl in doc.tables
                        for row in tbl.rows
                        for c in row.cells
                        if c.text.strip()
                    ]
                    all_text = " \n ".join(para_texts + tbl_texts)

                    # 1. No 'МОДУЛЬНЫЙ ДОКУМЕНТ'
                    self.assertNotIn(
                        "МОДУЛЬНЫЙ ДОКУМЕНТ",
                        all_text.upper(),
                        f"Template {tm.module_code} contains 'МОДУЛЬНЫЙ ДОКУМЕНТ'",
                    )

                    # 2. No upper approval stamp 'УТВЕРЖДЕНО' / 'УТВЕРЖДАЮ'
                    self.assertNotIn(
                        "УТВЕРЖДАЮ",
                        all_text.upper(),
                        f"Template {tm.module_code} upper stamp 'УТВЕРЖДАЮ' not deleted",
                    )

                    # 3. No footnote indicators like [1], [2], [15]
                    footnote_match = re.search(r"\[\d+\]", all_text)
                    self.assertIsNone(
                        footnote_match,
                        f"Template {tm.module_code} has leftover footnote: {footnote_match}",
                    )

                    # 4. No underscore sequences (____)
                    underscore_match = re.search(r"_{2,}", all_text)
                    self.assertIsNone(
                        underscore_match,
                        f"Template {tm.module_code} has leftover underscores: {underscore_match}",
                    )

                    # 5. No technical notes at the end
                    self.assertNotIn(
                        "Примечания по заполнению",
                        all_text,
                        f"Template {tm.module_code} notes section was not removed",
                    )
                    self.assertNotIn(
                        "Пояснения по заполнению",
                        all_text,
                        f"Template {tm.module_code} notes section was not removed",
                    )

                    # 6. Service zones MUST be preserved
                    self.assertIn(
                        "ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)",
                        all_text,
                        f"Template {tm.module_code} lost date/reg service zone",
                    )
                    self.assertIn(
                        "ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)",
                        all_text,
                        f"Template {tm.module_code} lost electronic signature visualization",
                    )

                    # 7. Signatory name substituted
                    self.assertIn(
                        "Петров П.П.",
                        all_text,
                        f"Template {tm.module_code} does not contain signatory name",
                    )

    async def test_project_card_creation_and_note_text(self):
        """Verify ProjectCard creation in DB, Section 2.3 note_text formatting, and approval route."""
        async with async_session_factory() as session:
            req = DocumentGenerateRequest(
                module_code="MOD-ORDER-IS",
                executor_name="Смирнов А.С.",
                executor_department="Управление информатизации",
                signatory_name="Кузнецов В.Н.",
                access_level="public",
                fields=SAMPLE_MODULE_FIELDS["MOD-ORDER-IS"],
                user_role="employee",
                project_title="Приказ о вводе в промышленную эксплуатацию АСУ ДОУ",
            )
            resp = await generate_document(session, req)

            # Retrieve ProjectCard from database
            res = await session.execute(
                select(ProjectCard).where(ProjectCard.id == uuid.UUID(resp.project_id))
            )
            card = res.scalars().first()
            self.assertIsNotNone(card, "ProjectCard must be found in DB")
            self.assertEqual(card.status, "draft")
            self.assertEqual(card.executor_name, "Смирнов А.С.")
            self.assertEqual(card.executor_department, "Управление информатизации")
            self.assertEqual(card.signatory_name, "Кузнецов В.Н.")
            self.assertEqual(card.project_title, "Приказ о вводе в промышленную эксплуатацию АСУ ДОУ")

            # Check Section 2.3 note_text format
            expected_prefix = "Проект подготовлен на основе модуля документа:"
            self.assertTrue(card.note_text.startswith(expected_prefix))
            self.assertIn("регистрационный № УТВ-01", card.note_text)
            self.assertIn("от 15.10.2026", card.note_text)
            self.assertIn("редакция (версия) 1.0", card.note_text)

            # Check attached sequential approval route
            self.assertTrue(len(card.approval_route) >= 4)
            stage_orders = [s["order"] for s in card.approval_route]
            self.assertEqual(stage_orders, sorted(stage_orders), "Stages must be strictly ordered")

            # Check Stage 1 has executor department
            self.assertIn("Управление информатизации", card.approval_route[0]["department"])

            # Check AuditLog
            audit_res = await session.execute(
                select(AuditLog).where(AuditLog.user_name == "Смирнов А.С.")
            )
            audit_entry = audit_res.scalars().first()
            self.assertIsNotNone(audit_entry, "AuditLog entry must be recorded")
            self.assertEqual(audit_entry.action_type, "generate_doc")

    async def test_fastapi_rest_endpoints(self):
        """Verify REST endpoints: /templates, /templates/{code}, /generate, /download, /projects."""
        transport = ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            # 1. GET /v1/documents/templates
            resp_list = await client.get("/v1/documents/templates")
            self.assertEqual(resp_list.status_code, 200)
            templates_data = resp_list.json()
            self.assertEqual(len(templates_data), 7)

            # 2. GET /v1/documents/templates/MOD-ORDER-IS
            resp_single = await client.get("/v1/documents/templates/MOD-ORDER-IS")
            self.assertEqual(resp_single.status_code, 200)
            single_data = resp_single.json()
            self.assertEqual(single_data["module_code"], "MOD-ORDER-IS")
            self.assertEqual(single_data["reg_number"], "УТВ-01")
            self.assertTrue(len(single_data["fields_schema"]) >= 5)

            # 3. GET /v1/documents/templates/УТВ-02 (by reg_number)
            resp_reg = await client.get("/v1/documents/templates/УТВ-02")
            self.assertEqual(resp_reg.status_code, 200)
            self.assertEqual(resp_reg.json()["module_code"], "MOD-ORDER-DOU")

            # 4. GET /v1/documents/templates/UNKNOWN (404)
            resp_404 = await client.get("/v1/documents/templates/NON-EXISTENT-CODE")
            self.assertEqual(resp_404.status_code, 404)

            # 5. POST /v1/documents/generate
            payload = {
                "module_code": "MOD-DISP-AUDIT",
                "executor_name": "Аудитор Тестовый",
                "executor_department": "Служба внутреннего контроля",
                "signatory_name": "Зампред Руководителя",
                "access_level": "public",
                "fields": SAMPLE_MODULE_FIELDS["MOD-DISP-AUDIT"],
                "user_role": "employee",
                "project_title": "Тестовое распоряжение об аудите",
            }
            resp_gen = await client.post("/v1/documents/generate", json=payload)
            self.assertEqual(resp_gen.status_code, 200)
            gen_data = resp_gen.json()
            self.assertIn("project_id", gen_data)
            self.assertIn("file_name", gen_data)
            self.assertIn("download_url", gen_data)
            self.assertTrue(len(gen_data["approval_route"]) >= 4)
            self.assertIn("Проект подготовлен на основе модуля", gen_data["note_text"])

            file_name = gen_data["file_name"]

            # 6. GET /v1/documents/download/{file_name}
            resp_dl = await client.get(f"/v1/documents/download/{file_name}")
            self.assertEqual(resp_dl.status_code, 200)
            self.assertEqual(
                resp_dl.headers.get("content-type"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            self.assertTrue(len(resp_dl.content) > 1000)

            # 7. GET /v1/documents/download/nonexistent.docx (404)
            resp_dl_404 = await client.get("/v1/documents/download/nonexistent_123.docx")
            self.assertEqual(resp_dl_404.status_code, 404)

            # 8. GET /v1/documents/projects
            resp_projects = await client.get("/v1/documents/projects")
            self.assertEqual(resp_projects.status_code, 200)
            projects_data = resp_projects.json()
            self.assertTrue(len(projects_data) >= 1)
            found = any(p["id"] == gen_data["project_id"] for p in projects_data)
            self.assertTrue(found, "Newly generated project card must appear in /v1/documents/projects")

    async def test_control_officer_default_phrase_handling(self):
        """Verify that default control phrase does not produce 'возложить на Контроль...'."""
        async with async_session_factory() as session:
            req = DocumentGenerateRequest(
                module_code="MOD-DISP-PROJECT",
                executor_name="Тестов Т.Т.",
                executor_department="Проектный департамент",
                signatory_name="Руководитель Р.Р.",
                access_level="public",
                fields={
                    "project_name": "Проект Ревизия",
                    "control_officer": "Контроль за исполнением настоящего распоряжения оставляю за собой",
                    "is_it_project": False,
                },
                user_role="employee",
            )
            resp = await generate_document(session, req)
            doc_path = settings.GENERATED_DOCS_DIR / resp.file_name
            doc = docx.Document(str(doc_path))
            all_text = " \n ".join(p.text for p in doc.paragraphs if p.text.strip())

            self.assertNotIn(
                "возложить на Контроль",
                all_text,
                "Must not create redundant 'возложить на Контроль...'",
            )
            self.assertIn(
                "Контроль за исполнением настоящего распоряжения оставляю за собой",
                all_text,
            )

    async def test_disp_project_it_conditional_routing(self):
        """Verify project committee without IT matches Rule 7 (no IT dept) and with IT matches Rule 8."""
        async with async_session_factory() as session:
            # 1. Non-IT committee (is_it_project=False) -> Rule 7
            req_no_it = DocumentGenerateRequest(
                module_code="MOD-DISP-PROJECT",
                executor_name="Иванов И.И.",
                executor_department="Финансовый департамент",
                signatory_name="Петров П.П.",
                fields={"project_name": "Бюджетный комитет", "is_it_project": False},
            )
            resp_no_it = await generate_document(session, req_no_it)
            depts_no_it = [s.department.lower() for s in resp_no_it.approval_route]
            self.assertFalse(
                any("информационных технологий" in d for d in depts_no_it),
                "Non-IT committee must NOT include IT Department in approval route",
            )

            # 2. IT committee (is_it_project=True) -> Rule 8
            req_it = DocumentGenerateRequest(
                module_code="MOD-DISP-PROJECT",
                executor_name="Иванов И.И.",
                executor_department="Департамент ИТ",
                signatory_name="Петров П.П.",
                fields={"project_name": "Комитет АСУ", "is_it_project": True},
            )
            resp_it = await generate_document(session, req_it)
            depts_it = [s.department.lower() for s in resp_it.approval_route]
            self.assertTrue(
                any("информационных технологий" in d for d in depts_it),
                "IT committee MUST include IT Department in approval route",
            )

    async def test_order_dou_paragraph_numbering_with_and_without_repeal(self):
        """Verify paragraph numbering in MOD-ORDER-DOU: repeal is 5, control is 6; without repeal, control is 5."""
        async with async_session_factory() as session:
            # With repeal
            req_rep = DocumentGenerateRequest(
                module_code="MOD-ORDER-DOU",
                executor_name="Иванов И.И.",
                executor_department="Департамент ДОУ",
                signatory_name="Петров П.П.",
                fields={"repeal_order": "от 10.01.2024 № ОД-1"},
            )
            resp_rep = await generate_document(session, req_rep)
            doc_rep = docx.Document(str(settings.GENERATED_DOCS_DIR / resp_rep.file_name))
            rep_texts = [p.text for p in doc_rep.paragraphs if p.text.strip()]
            self.assertTrue(any(t.startswith("5. Признать утратившим силу") for t in rep_texts))
            self.assertTrue(any(t.startswith("6. Контроль за исполнением") for t in rep_texts))

            # Without repeal
            req_norep = DocumentGenerateRequest(
                module_code="MOD-ORDER-DOU",
                executor_name="Иванов И.И.",
                executor_department="Департамент ДОУ",
                signatory_name="Петров П.П.",
                fields={},
            )
            resp_norep = await generate_document(session, req_norep)
            doc_norep = docx.Document(str(settings.GENERATED_DOCS_DIR / resp_norep.file_name))
            norep_texts = [p.text for p in doc_norep.paragraphs if p.text.strip()]
            self.assertFalse(any("Признать утратившим силу" in t for t in norep_texts))
            self.assertTrue(any(t.startswith("5. Контроль за исполнением") for t in norep_texts))


if __name__ == "__main__":
    unittest.main()
