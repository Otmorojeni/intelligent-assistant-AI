"""Interactive DOCX Document Generator for modular orders and dispositions.

Implements document generation, field substitution, and strict cleaning according
to Section 3.3 and Project Card creation per Section 2.3 of
'Методические_рекомендации_по_работе_с_модульными_документами.docx'.
"""

from datetime import datetime, timezone
import logging
from pathlib import Path
import re
import unicodedata
from typing import Any, Dict, List, Optional, Tuple
import uuid

import docx
from docx import Document
from fastapi import HTTPException, status
from sqlalchemy import select, desc
from sqlalchemy.orm import joinedload
from sqlalchemy.ext.asyncio import AsyncSession


from src.config import settings
from src.core.approval_router import calculate_approval_route
from src.db.models import ApprovalRule, AuditLog, ProjectCard, TemplateModule, utc_now
from src.schemas.api_models import (
    ApprovalCalculateRequest,
    ApprovalStage,
    DocumentGenerateRequest,
    DocumentGenerateResponse,
    ProjectCardItem,
    TemplateModuleInfo,
)

logger = logging.getLogger(__name__)


def is_truthy(val: Any) -> bool:
    """Robust boolean parser supporting bool, int, and textual representations."""
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return val != 0
    if isinstance(val, str):
        return val.strip().lower() in ("true", "1", "yes", "да", "y", "t")
    return bool(val)


def find_template_file(materials_dir: Path, file_name: str) -> Path:
    """Locate template file in materials directory with Unicode NFC normalization support."""
    direct = materials_dir / file_name
    if direct.exists():
        return direct

    target_nfc = unicodedata.normalize("NFC", file_name).lower()
    for p in materials_dir.glob("*.docx"):
        if unicodedata.normalize("NFC", p.name).lower() == target_nfc:
            return p

    raise FileNotFoundError(f"Template file '{file_name}' not found in {materials_dir}")


def remove_approval_stamp(doc: Document) -> None:
    """Delete upper approval stamp table (e.g., 'УТВЕРЖДАЮ...', Table 0) per Section 3.3."""
    for tbl in list(doc.tables):
        txt = " ".join(c.text for r in tbl.rows for c in r.cells).upper()
        if "УТВЕРЖДАЮ" in txt or "УТВЕРЖДЕНО" in txt or "ГРИФ УТВЕРЖДЕНИЯ" in txt:
            tbl._element.getparent().remove(tbl._element)

    # Also clean any top paragraphs before document heading with approval stamp
    for p in list(doc.paragraphs[:5]):
        txt = p.text.strip().upper()
        if "УТВЕРЖДАЮ" in txt or "УТВЕРЖДЕНО" in txt or "ГРИФ УТВЕРЖДЕНИЯ" in txt:
            p._element.getparent().remove(p._element)


def remove_modular_document_label(doc: Document) -> None:
    """Delete label 'МОДУЛЬНЫЙ ДОКУМЕНТ' per Section 3.3."""
    for p in list(doc.paragraphs):
        txt = p.text.strip().upper()
        if "МОДУЛЬНЫЙ ДОКУМЕНТ" in txt:
            p._element.getparent().remove(p._element)


def remove_notes_and_explanations(doc: Document) -> None:
    """Delete technical notes, explanations, and instructions at the end of the document."""
    in_notes = False
    for el in list(doc._body._element):
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p_obj = docx.text.paragraph.Paragraph(el, doc)
            txt = p_obj.text.strip().lower()
            if (
                txt.startswith("примечания")
                or txt.startswith("пояснения")
                or "примечания по заполнению" in txt
                or "пояснения по заполнению" in txt
            ):
                in_notes = True
        elif tag == "sectPr":
            break

        if in_notes:
            el.getparent().remove(el)


def clean_text_strict(text: str, signatory_name: str = "") -> str:
    """Apply strict cleaning rules per Section 3.3 while preserving service zones."""
    if not text:
        return ""

    # Preserve service zones marked with 'НЕ УДАЛЯТЬ'
    if "НЕ УДАЛЯТЬ" in text or "ВИЗУАЛИЗАЦИЯ ЭП" in text:
        return text

    res = text

    # Signatory substitution
    if signatory_name:
        res = re.sub(r"И\.\s*О\.\s*Фамилия", signatory_name, res)

    # Delete footnote numbers [1], [2], [15], etc.
    res = re.sub(r"\[\s*\d+(\s*,\s*\d+)*\s*\]", "", res)

    # Delete unused underscore lines (____, _____)
    res = re.sub(r"_{2,}", "", res)
    res = re.sub(r"(?:_\s*){3,}", "", res)

    # Delete empty brackets [] or [ ]
    res = re.sub(r"\[\s*\]", "", res)

    # Clean unused option slashes and variants
    res = re.sub(r"\[([^\]]*)\]", r"\1", res)
    res = res.replace("разработанного / измененного", "разработанного")
    res = res.replace("процесса / процессов", "процесса")
    res = res.replace("проекту / проектам", "проекту")
    res = res.replace("проекта / проектов", "проекта")
    res = res.replace("реализуемого(ых)", "реализуемого")

    # Clean whitespace and punctuation spaces
    res = re.sub(r"[ \t]+", " ", res)
    res = re.sub(r"\s+([,.:;?!])", r"\1", res)
    res = re.sub(r"([«“])\s+", r"\1", res)
    res = re.sub(r"\s+([»”])", r"\1", res)

    return res.strip()


def clean_empty_paragraphs(doc: Document) -> None:
    """Clean redundant blank paragraphs left after removing headers/stamps."""
    # Strip leading empty paragraphs
    while doc.paragraphs and not doc.paragraphs[0].text.strip():
        doc.paragraphs[0]._element.getparent().remove(doc.paragraphs[0]._element)

    # Strip trailing empty paragraphs
    while doc.paragraphs and not doc.paragraphs[-1].text.strip():
        doc.paragraphs[-1]._element.getparent().remove(doc.paragraphs[-1]._element)

    # Remove consecutive empty paragraphs
    prev_empty = False
    for p in list(doc.paragraphs):
        is_empty = not p.text.strip()
        if is_empty and prev_empty:
            p._element.getparent().remove(p._element)
        prev_empty = is_empty


def substitute_template_fields(
    doc: Document,
    module_code: str,
    fields: Dict[str, Any],
    signatory_name: str,
    executor_name: str,
    executor_department: str,
) -> None:
    """Perform module-specific substitutions and general cleaning across paragraphs and tables."""
    code = module_code.upper()

    # 1. Module-specific paragraph replacements
    if code == "MOD-ORDER-IS":
        system_name = fields.get("system_name", "Информационная система")
        purpose = fields.get("purpose", "автоматизации процессов обработки информации")
        commissioning_date = fields.get("commissioning_date", datetime.now().strftime("%d.%m.%Y"))
        responsible_dept = fields.get("responsible_dept", executor_department or "Департамент информационных технологий")
        infosec_dept = fields.get("infosec_dept", "Департамент информационной безопасности")
        control_officer = fields.get("control_officer", "директор Департамента информационных технологий")

        for p in doc.paragraphs:
            txt = p.text
            if "О вводе в эксплуатацию" in txt and ("____" in txt or "[" in txt):
                align = p.alignment
                p.text = f"О вводе в эксплуатацию информационной системы «{system_name}»"
                p.alignment = align
            elif "содержащего решение о готовности к вводу в эксплуатацию" in txt:
                align = p.alignment
                p.text = (
                    f"На основании решения комиссии о готовности к вводу в эксплуатацию компонента "
                    f"информационной системы «{system_name}» (далее — «{system_name}»), разработанного "
                    f"в соответствии с утвержденным техническим заданием,"
                )
                p.alignment = align
            elif txt.strip().startswith("1. Ввести"):
                align = p.alignment
                p.text = f"1. Ввести «{system_name}» в эксплуатацию с {commissioning_date}."
                p.alignment = align
            elif "обеспечить применение" in txt:
                align = p.alignment
                p.text = f"обеспечить применение «{system_name}» по следующему назначению: {purpose}."
                p.alignment = align
            elif "3.1. функционирование" in txt:
                align = p.alignment
                p.text = f"3.1. функционирование «{system_name}» в соответствии с эксплуатационной документацией;"
                p.alignment = align
            elif "3.2. сопровождение" in txt:
                align = p.alignment
                p.text = f"3.2. сопровождение «{system_name}» с привлечением при необходимости внешних подрядных организаций."
                p.alignment = align
            elif "4. Ответственным за исполнение настоящего приказа" in txt:
                align = p.alignment
                if "оставляю за собой" in control_officer.lower() or control_officer.lower().startswith("контроль"):
                    p.text = "4. Контроль за исполнением настоящего приказа оставляю за собой."
                else:
                    p.text = f"4. Ответственным за исполнение настоящего приказа является {control_officer}."
                p.alignment = align
            elif txt.strip().startswith("2.") and "____" in txt:
                align = p.alignment
                p.text = f"2. {responsible_dept}:"
                p.alignment = align

    elif code == "MOD-ORDER-DOU":
        dept = fields.get("department_name", executor_department or "Структурное подразделение")
        primary = fields.get("primary_responsible", executor_name or "ответственный работник")
        deputy = fields.get("deputy_responsible", "замещающий работник")
        reason = fields.get("reason", "Инструкции по делопроизводству № ОД-184")
        control = fields.get("control_officer", "руководителя Департамента ДОУ")

        for p in list(doc.paragraphs):
            txt = p.text
            if "В целях организации делопроизводства" in txt:
                align = p.alignment
                p.text = f"В целях организации делопроизводства и соблюдения требований {reason},"
                p.alignment = align
            elif "3. Руководителям подразделений обеспечивать замещение" in txt:
                align = p.alignment
                p.text = "3. Руководителям подразделений обеспечивать замещение ответственных лиц на период отсутствия и сообщать об изменениях в Департамент ДОУ."
                p.alignment = align
            elif "осуществлять методическое сопровождение" in txt:
                align = p.alignment
                p.text = "4. Департаменту ДОУ осуществлять методическое сопровождение и координацию работы ответственных лиц."
                p.alignment = align
            elif "Признать утратившим силу" in txt:
                repeal = fields.get("repeal_order", "")
                if repeal:
                    align = p.alignment
                    p.text = f"5. Признать утратившим силу приказ {repeal}."
                    p.alignment = align
                else:
                    p._element.getparent().remove(p._element)
            elif "Контроль за исполнением настоящего приказа" in txt:
                repeal = fields.get("repeal_order", "")
                num = "6" if repeal else "5"
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = f"{num}. Контроль за исполнением настоящего приказа оставляю за собой."
                else:
                    p.text = f"{num}. Контроль за исполнением настоящего приказа возложить на {control}."
                p.alignment = align

        # Fill Table 2 (appointed staff)
        if len(doc.tables) >= 2:
            tbl = doc.tables[1]  # table after stamp was removed
            if len(tbl.rows) >= 3:
                tbl.rows[1].cells[0].text = dept
                tbl.rows[1].cells[1].text = f"{primary} (основной ответственный)"
                tbl.rows[2].cells[0].text = dept
                tbl.rows[2].cells[1].text = f"{deputy} (замещающий ответственный)"
                # Remove extra template placeholder rows if any
                while len(tbl.rows) > 3:
                    tbl._element.remove(tbl.rows[-1]._element)

    elif code == "MOD-ORDER-PROP":
        dept = fields.get("department_name", executor_department or "Структурное подразделение")
        transferor = fields.get("transferor_name", "сдающее лицо")
        transferee = fields.get("transferee_name", "принимающее лицо")
        basis = fields.get("transfer_basis", "кадровыми изменениями")
        deadline = fields.get("deadline", datetime.now().strftime("%d.%m.%Y"))
        commission = fields.get("commission_members", "главного специалиста аппарата")
        control = fields.get("control_officer", "руководителя аппарата")

        for p in doc.paragraphs:
            txt = p.text
            if "В связи с" in txt and "____" in txt:
                align = p.alignment
                p.text = f"В связи с {basis},"
                p.alignment = align
            elif txt.strip().startswith("1.") and ("сдать" in txt or "принять" in txt):
                align = p.alignment
                p.text = (
                    f"1. {transferor} сдать, а {transferee} принять дела и имущество подразделения {dept} "
                    f"в срок до {deadline} в порядке, предусмотренном Положением о порядке приема и передачи "
                    f"дел и имущества при назначении на должность и освобождении от должности руководителей "
                    f"структурных подразделений Организации."
                )
                p.alignment = align
            elif "2. Назначить" in txt and "представителем" in txt:
                align = p.alignment
                if "представител" in commission.lower():
                    p.text = f"2. Назначить {commission} при приеме и передаче дел и имущества."
                else:
                    p.text = f"2. Назначить {commission} представителем Организации при приеме и передаче дел и имущества."
                p.alignment = align
            elif "3. Контроль за исполнением настоящего приказа" in txt:
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = "3. Контроль за исполнением настоящего приказа оставляю за собой."
                else:
                    p.text = f"3. Контроль за исполнением настоящего приказа возложить на {control}."
                p.alignment = align

    elif code == "MOD-ORDER-LNA":
        act_title = fields.get("act_title", "Локальный нормативный акт")
        purpose = fields.get("purpose", "совершенствования внутренних процессов Организации")
        effective_date = fields.get("effective_date", datetime.now().strftime("%d.%m.%Y"))
        resp_dept = fields.get("responsible_dept", executor_department or "Департамент ДОУ")
        placement = fields.get("placement_location", "на корпоративном портале Организации")
        repeal_act = fields.get("repeal_act", "")
        control = fields.get("control_officer", "Заместителя Председателя Организации")

        for p in list(doc.paragraphs):
            txt = p.text
            if "Об утверждении и введении в действие" in txt:
                align = p.alignment
                p.text = f"Об утверждении и введении в действие документа «{act_title}»"
                p.alignment = align
            elif "В целях" in txt and "____" in txt:
                align = p.alignment
                p.text = f"В целях {purpose},"
                p.alignment = align
            elif txt.strip().startswith("1. Утвердить"):
                align = p.alignment
                p.text = f"1. Утвердить {act_title} согласно приложению к настоящему приказу."
                p.alignment = align
            elif "2. Ввести указанный документ в действие" in txt:
                align = p.alignment
                p.text = f"2. Ввести указанный документ в действие с {effective_date}."
                p.alignment = align
            elif txt.strip().startswith("3.") and "обеспечить:" in txt:
                align = p.alignment
                p.text = f"3. {resp_dept} обеспечить:"
                p.alignment = align
            elif "размещение актуальной редакции документа" in txt:
                align = p.alignment
                p.text = f"• размещение актуальной редакции документа {placement}."
                p.alignment = align
            elif "Признать утратившим силу" in txt:
                if repeal_act:
                    align = p.alignment
                    p.text = f"4. Признать утратившим силу {repeal_act}."
                    p.alignment = align
                else:
                    p._element.getparent().remove(p._element)
            elif "Контроль за исполнением настоящего приказа" in txt:
                num = "5" if repeal_act else "4"
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = f"{num}. Контроль за исполнением настоящего приказа оставляю за собой."
                else:
                    p.text = f"{num}. Контроль за исполнением настоящего приказа возложить на {control}."
                p.alignment = align

    elif code == "MOD-DISP-AUDIT":
        process = fields.get("audit_process", "процессов Организации")
        goal = fields.get("audit_goal", "оценки результативности процессов и управления связанными рисками")
        period = fields.get("audit_period", "в течение 14 календарных дней")
        audited_dept = fields.get("audited_dept", executor_department or "Проверяемое подразделение")
        leader = fields.get("audit_leader", executor_name or "главный аудитор")
        team = fields.get("audit_team", "аудиторская группа")
        deadline = fields.get("report_deadline", datetime.now().strftime("%d.%m.%Y"))
        control = fields.get("control_officer", "директора Департамента внутреннего аудита")

        for p in doc.paragraphs:
            txt = p.text
            if "О проведении внутреннего аудита" in txt:
                align = p.alignment
                p.text = f"О проведении внутреннего аудита процесса {process}"
                p.alignment = align
            elif "В соответствии с" in txt and "в целях оценки" in txt:
                align = p.alignment
                p.text = f"В соответствии с планом внутреннего аудита, в целях {goal},"
                p.alignment = align
            elif txt.strip().startswith("1. Провести"):
                align = p.alignment
                p.text = f"1. Провести {period} внутренний аудит процесса {process}."
                p.alignment = align
            elif "обеспечить аудиторской группе доступ" in txt:
                align = p.alignment
                p.text = f"4. {audited_dept} обеспечить аудиторской группе доступ к документам, информационным системам и работникам в пределах программы аудита."
                p.alignment = align
            elif "отчет о результатах аудита" in txt:
                align = p.alignment
                p.text = f"5. Руководителю аудиторской группы до {deadline} представить руководству отчет о результатах аудита."
                p.alignment = align
            elif "подготовить план мероприятий" in txt:
                align = p.alignment
                p.text = f"6. {audited_dept} подготовить план мероприятий по результатам аудита."
                p.alignment = align
            elif "7. Контроль за исполнением настоящего распоряжения" in txt:
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = "7. Контроль за исполнением настоящего распоряжения оставляю за собой."
                else:
                    p.text = f"7. Контроль за исполнением настоящего распоряжения возложить на {control}."
                p.alignment = align

        # Fill Table 2 (audit team)
        if len(doc.tables) >= 2:
            tbl = doc.tables[1]
            if len(tbl.rows) >= 3:
                tbl.rows[1].cells[1].text = leader
                tbl.rows[2].cells[1].text = team

    elif code == "MOD-DISP-PROJECT":
        p_name = fields.get("project_name", "Проект развития")
        purpose = fields.get("committee_purpose", "обеспечения координации и контроля реализации проекта")
        reason = fields.get("reason", "планом реализации стратегических инициатив")
        dept = fields.get("managing_dept", executor_department or "Проектный офис")
        chairman = fields.get("chairman", "Заместитель Председателя Организации")
        deputy_ch = fields.get("deputy_chairman", "Директор департамента")
        secretary = fields.get("secretary", "руководитель проектов")
        members = fields.get("members", "руководители заинтересованных подразделений")
        repeal = fields.get("repeal_order", "")
        control = fields.get("control_officer", "Председателя Организации")

        for p in list(doc.paragraphs):
            txt = p.text
            if "О создании проектного комитета" in txt:
                align = p.alignment
                rep_txt = f" и об отмене распоряжения {repeal}" if repeal else ""
                p.text = f"О создании проектного комитета по проекту «{p_name}»{rep_txt}"
                p.alignment = align
            elif "В целях" in txt and "для обеспечения координации" in txt:
                align = p.alignment
                p.text = (
                    f"В целях {purpose}, в связи с {reason}, для обеспечения координации и контроля "
                    f"реализации проекта «{p_name}», реализуемого подразделением {dept},"
                )
                p.alignment = align
            elif txt.strip().startswith("1. Создать Проектный комитет"):
                align = p.alignment
                p.text = f"1. Создать Проектный комитет по проекту «{p_name}» (далее — Комитет) в составе:"
                p.alignment = align
            elif "Отменить распоряжение Организации" in txt:
                if repeal:
                    align = p.alignment
                    p.text = f"2. Отменить распоряжение Организации {repeal}."
                    p.alignment = align
                else:
                    p._element.getparent().remove(p._element)
            elif "Контроль за исполнением настоящего распоряжения" in txt:
                num = "3" if repeal else "2"
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = f"{num}. Контроль за исполнением настоящего распоряжения оставляю за собой."
                else:
                    p.text = f"{num}. Контроль за исполнением настоящего распоряжения возложить на {control}."
                p.alignment = align

        # Fill Table 2 (committee members)
        if len(doc.tables) >= 2:
            tbl = doc.tables[1]
            if len(tbl.rows) >= 5:
                tbl.rows[1].cells[1].text = chairman
                tbl.rows[2].cells[1].text = deputy_ch or "Не назначается"
                tbl.rows[3].cells[1].text = secretary
                tbl.rows[4].cells[1].text = members

        # Ensure signatory block is present before ВИЗУАЛИЗАЦИЯ ЭП
        for idx, p in enumerate(doc.paragraphs):
            if "ВИЗУАЛИЗАЦИЯ ЭП" in p.text:
                prev_p = doc.paragraphs[idx - 1] if idx > 0 else None
                if not prev_p or signatory_name not in prev_p.text:
                    p.insert_paragraph_before(f"Заместитель Председателя Организации        {signatory_name}")
                break

    elif code == "MOD-DISP-WORKGROUP":

        subj = fields.get("workgroup_subject", "подготовке предложений")
        goal = fields.get("goal", "выработки согласованных решений")
        leader = fields.get("leader", executor_name or "руководитель рабочей группы")
        secretary = fields.get("secretary", "секретарь рабочей группы")
        members = fields.get("members", "члены рабочей группы")
        deadline = fields.get("deadline", datetime.now().strftime("%d.%m.%Y"))
        result = fields.get("expected_result", "итоговый отчет и предложения")
        termination = fields.get("termination_event", "представления итогового отчета")
        control = fields.get("control_officer", "Заместителя Председателя Организации")

        for p in doc.paragraphs:
            txt = p.text
            if "О создании рабочей группы по" in txt:
                align = p.alignment
                p.text = f"О создании рабочей группы по {subj}"
                p.alignment = align
            elif "В целях" in txt and "____" in txt:
                align = p.alignment
                p.text = f"В целях {goal},"
                p.alignment = align
            elif txt.strip().startswith("1. Создать рабочую группу"):
                align = p.alignment
                p.text = f"1. Создать рабочую группу по {subj} (далее — Рабочая группа) в составе:"
                p.alignment = align
            elif txt.strip().startswith("2. Рабочей группе до"):
                align = p.alignment
                p.text = f"2. Рабочей группе до {deadline} выполнить возложенные задачи и представить {result}."
                p.alignment = align
            elif "5. Рабочая группа прекращает деятельность" in txt:
                align = p.alignment
                p.text = f"5. Рабочая группа прекращает деятельность после {termination}."
                p.alignment = align
            elif "6. Контроль за исполнением настоящего распоряжения" in txt:
                align = p.alignment
                if "оставляю за собой" in control.lower() or control.lower().startswith("контроль"):
                    p.text = "6. Контроль за исполнением настоящего распоряжения оставляю за собой."
                else:
                    p.text = f"6. Контроль за исполнением настоящего распоряжения возложить на {control}."
                p.alignment = align

        # Fill Table 2 (roles)
        if len(doc.tables) >= 2:
            tbl = doc.tables[1]
            if len(tbl.rows) >= 4:
                tbl.rows[1].cells[1].text = leader
                tbl.rows[2].cells[1].text = secretary
                tbl.rows[3].cells[1].text = members

    # 2. General cleaning pass for all remaining paragraphs
    for p in list(doc.paragraphs):
        txt = p.text
        if not txt.strip():
            continue
        cleaned = clean_text_strict(txt, signatory_name)
        if cleaned != txt:
            align = p.alignment
            p.text = cleaned
            p.alignment = align

    # 3. General cleaning pass for all table cells
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    txt = p.text
                    if not txt.strip():
                        continue
                    cleaned = clean_text_strict(txt, signatory_name)
                    if cleaned != txt:
                        align = p.alignment
                        p.text = cleaned
                        p.alignment = align


def clean_and_populate_docx(
    template_path: Path,
    module_code: str,
    fields: Dict[str, Any],
    signatory_name: str,
    executor_name: str,
    executor_department: str,
    output_path: Path,
) -> None:
    """Execute end-to-end loading, cleaning per Section 3.3, and saving of generated DOCX."""
    doc = Document(str(template_path))

    # 1. Delete upper approval stamp table (Table 0)
    remove_approval_stamp(doc)

    # 2. Delete label 'МОДУЛЬНЫЙ ДОКУМЕНТ'
    remove_modular_document_label(doc)

    # 3. Delete notes and technical explanations at the end
    remove_notes_and_explanations(doc)

    # 4. Substitute fields & apply strict Section 3.3 text cleaning
    substitute_template_fields(
        doc=doc,
        module_code=module_code,
        fields=fields,
        signatory_name=signatory_name,
        executor_name=executor_name,
        executor_department=executor_department,
    )

    # 5. Clean empty paragraphs left after deletions
    clean_empty_paragraphs(doc)

    # Save to output file
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    logger.info(f"Generated clean DOCX saved to {output_path}")


async def get_available_templates(session: AsyncSession) -> List[TemplateModuleInfo]:
    """Retrieve all available template modules with metadata and dynamic field schemas."""
    res = await session.execute(select(TemplateModule).order_by(TemplateModule.reg_number))
    modules = res.scalars().all()
    items: List[TemplateModuleInfo] = []
    for m in modules:
        items.append(
            TemplateModuleInfo(
                module_code=m.module_code,
                doc_kind=m.doc_kind,
                title=m.title,
                reg_number=m.reg_number,
                version=m.version,
                file_name=m.file_name,
                fields_schema=m.fields_schema if isinstance(m.fields_schema, list) else [],
            )
        )
    return items


async def get_template_by_code(
    session: AsyncSession,
    module_code: str,
) -> Optional[TemplateModuleInfo]:
    """Find specific template module by code or registration number (e.g., MOD-ORDER-IS or УТВ-01)."""
    clean_code = module_code.strip()
    res = await session.execute(
        select(TemplateModule).where(
            (TemplateModule.module_code.ilike(clean_code))
            | (TemplateModule.reg_number.ilike(clean_code))
        )
    )
    m = res.scalars().first()
    if not m:
        return None

    return TemplateModuleInfo(
        module_code=m.module_code,
        doc_kind=m.doc_kind,
        title=m.title,
        reg_number=m.reg_number,
        version=m.version,
        file_name=m.file_name,
        fields_schema=m.fields_schema if isinstance(m.fields_schema, list) else [],
    )


async def get_all_projects(session: AsyncSession) -> List[ProjectCardItem]:
    """Retrieve all created project cards ordered by creation date descending."""
    res = await session.execute(
        select(ProjectCard)
        .options(joinedload(ProjectCard.template))
        .order_by(desc(ProjectCard.created_at))
    )

    cards = res.scalars().unique().all()
    items: List[ProjectCardItem] = []
    for c in cards:
        items.append(
            ProjectCardItem(
                id=str(c.id),
                template_id=str(c.template_id),
                module_code=c.template.module_code if c.template else None,
                project_title=c.project_title,
                executor_name=c.executor_name,
                executor_department=c.executor_department,
                signatory_name=c.signatory_name,
                access_level=c.access_level,
                note_text=c.note_text,
                filled_fields=c.filled_fields if isinstance(c.filled_fields, dict) else {},
                approval_route=[ApprovalStage(**s) for s in c.approval_route] if isinstance(c.approval_route, list) else [],
                generated_file_name=c.generated_file_name,
                download_url=f"/v1/documents/download/{c.generated_file_name}",
                status=c.status,
                created_at=c.created_at,
            )
        )
    return items


async def generate_document(
    session: AsyncSession,
    request: DocumentGenerateRequest,
) -> DocumentGenerateResponse:
    """Generate clean DOCX document, calculate approval route, and register Project Card in DB."""
    t0 = datetime.now()

    # 1. Fetch template module from DB
    clean_code = request.module_code.strip()
    res = await session.execute(
        select(TemplateModule).where(
            (TemplateModule.module_code.ilike(clean_code))
            | (TemplateModule.reg_number.ilike(clean_code))
        )
    )
    template = res.scalars().first()
    if not template:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Template module '{request.module_code}' not found in registry.",
        )

    # 2. Find template file in materials
    template_path = find_template_file(settings.MATERIALS_DIR, template.file_name)

    # 3. Generate unique file name
    unique_suffix = uuid.uuid4().hex[:10]
    out_file_name = f"{template.module_code.lower()}_{unique_suffix}.docx"
    output_path = settings.GENERATED_DOCS_DIR / out_file_name

    # 4. Execute cleaning & substitution per Section 3.3
    clean_and_populate_docx(
        template_path=template_path,
        module_code=template.module_code,
        fields=request.fields,
        signatory_name=request.signatory_name,
        executor_name=request.executor_name,
        executor_department=request.executor_department,
        output_path=output_path,
    )

    # 5. Calculate sequential approval route via approval_router.py
    infosec_access = is_truthy(request.fields.get("affects_is_access", False)) or (template.module_code == "MOD-ORDER-IS")
    it_involved = is_truthy(request.fields.get("is_it_project", False)) or (template.module_code == "MOD-ORDER-IS")
    regional_involved = is_truthy(request.fields.get("involves_regional", False))

    calc_req = ApprovalCalculateRequest(
        doc_type=f"{template.module_code} {template.title}",
        uses_module=True,
        involves_infosec_access=infosec_access,
        involves_it=it_involved,
        involves_regional=regional_involved,
        initiator_department=request.executor_department,
    )
    route_resp = await calculate_approval_route(session, calc_req)

    # 6. Format note_text strictly according to Section 2.3 of Methodological Guidelines
    note_date = request.fields.get("module_date", "15.10.2026")
    note_text = (
        f"Проект подготовлен на основе модуля документа: {template.title}, "
        f"регистрационный № {template.reg_number} от {note_date}, редакция (версия) {template.version}"
    )

    # 7. Determine project title
    if request.project_title:
        project_title = request.project_title
    elif "system_name" in request.fields:
        project_title = f"{template.title} «{request.fields['system_name']}»"
    elif "act_title" in request.fields:
        project_title = f"{template.title} «{request.fields['act_title']}»"
    elif "project_name" in request.fields:
        project_title = f"{template.title} по проекту «{request.fields['project_name']}»"
    elif "workgroup_subject" in request.fields:
        project_title = f"{template.title} по {request.fields['workgroup_subject']}"
    elif "audit_process" in request.fields:
        project_title = f"{template.title} процесса {request.fields['audit_process']}"
    elif "department_name" in request.fields:
        project_title = f"{template.title} подразделения {request.fields['department_name']}"
    else:
        project_title = template.title

    # 8. Create ProjectCard in DB
    project_card = ProjectCard(
        id=uuid.uuid4(),
        template_id=template.id,
        project_title=project_title,
        executor_name=request.executor_name,
        executor_department=request.executor_department,
        signatory_name=request.signatory_name,
        access_level=request.access_level or "public",
        note_text=note_text,
        filled_fields=request.fields,
        approval_route=[s.model_dump() for s in route_resp.stages],
        generated_file_name=out_file_name,
        status="draft",
        created_at=utc_now(),
    )
    session.add(project_card)

    # 9. Audit trail logging
    latency_ms = (datetime.now() - t0).total_seconds() * 1000.0
    audit = AuditLog(
        id=uuid.uuid4(),
        timestamp=utc_now(),
        user_role=request.user_role or "employee",
        user_name=request.executor_name,
        action_type="generate_doc",
        query_text=f"Generate {template.module_code}: {project_title}",
        response_summary=f"Project {project_card.id} generated file {out_file_name}",
        status="success",
        latency_ms=latency_ms,
    )
    session.add(audit)
    await session.commit()
    await session.refresh(project_card)

    logger.info(f"ProjectCard created successfully: ID={project_card.id}, file={out_file_name}")

    return DocumentGenerateResponse(
        project_id=str(project_card.id),
        file_name=out_file_name,
        download_url=f"/v1/documents/download/{out_file_name}",
        note_text=note_text,
        approval_route=route_resp.stages,
        created_at=project_card.created_at,
    )
