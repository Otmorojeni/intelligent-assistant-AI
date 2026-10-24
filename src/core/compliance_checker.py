"""Normative Compliance Checker for Draft DOCX Documents (OD-184 & Methodological Guidelines).

Inspects uploaded .docx draft orders, dispositions, and internal documents against 16 mandatory
regulatory formatting and structure rules from:
- Инструкция по ДОУ № ОД-184 (разделы 5, 6, 9, 10)
- Методические рекомендации по работе с модульными документами (разделы 2.3, 3.3, 4.2)

Provides automated document remediation (Auto-Fix) generating clean, compliant .docx files.
"""

import io
import logging
from pathlib import Path
import re
from typing import List, Optional, Tuple
import uuid

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt
from sqlalchemy.ext.asyncio import AsyncSession

from src.config import settings
from src.core.approval_router import calculate_approval_route
from src.schemas.api_models import (
    ApprovalCalculateRequest,
    ApprovalRecommendation,
    ComplianceCheckResponse,
    ComplianceIssue,
)

logger = logging.getLogger(__name__)

ALLOWED_FONTS = {"times new roman", "arial", "pt astra serif"}


def _extract_docx_lines_and_tables(doc: docx.Document) -> Tuple[List[str], List[str]]:
    """Extract paragraph lines and table cell texts from docx.Document."""
    para_lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    table_cells: List[str] = []
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    table_cells.append(txt)
    return para_lines, table_cells


def _extract_doc_defaults(doc: docx.Document) -> Tuple[Optional[str], Optional[float]]:
    """Extract default font family and size in pt from document styles (w:docDefaults)."""
    default_font: Optional[str] = None
    default_size_pt: Optional[float] = None
    try:
        el = doc.styles.element
        rpr_defaults = el.xpath("./w:docDefaults/w:rPrDefault/w:rPr")
        if rpr_defaults:
            rpr = rpr_defaults[0]
            rfonts = rpr.find(docx.oxml.ns.qn("w:rFonts"))
            if rfonts is not None:
                default_font = (
                    rfonts.get(docx.oxml.ns.qn("w:ascii"))
                    or rfonts.get(docx.oxml.ns.qn("w:hAnsi"))
                    or rfonts.get(docx.oxml.ns.qn("w:cs"))
                )
                if not default_font:
                    theme = rfonts.get(docx.oxml.ns.qn("w:asciiTheme"))
                    if theme and "minor" in theme.lower():
                        default_font = "Calibri"
            sz = rpr.find(docx.oxml.ns.qn("w:sz"))
            if sz is not None and sz.get(docx.oxml.ns.qn("w:val")):
                try:
                    default_size_pt = float(sz.get(docx.oxml.ns.qn("w:val"))) / 2.0
                except (ValueError, TypeError):
                    pass
    except Exception:
        pass
    return default_font, default_size_pt


def _detect_doc_kind_and_module(
    all_text: str,
    doc: Optional[docx.Document] = None,
) -> Tuple[str, Optional[str]]:
    """Detect document kind and matching modular template code from text using strict structural patterns."""
    up = all_text.upper()
    doc_kind = "Неопознанный документ / Не является ОРД"

    # 1. Structural check on header paragraphs (Реквизит 09: Наименование вида документа)
    candidate_paras: List[str] = []
    if doc and doc.paragraphs:
        candidate_paras = [p.text.strip() for p in doc.paragraphs[:15] if p.text.strip()]
    else:
        candidate_paras = [line.strip() for line in all_text.split("\n")[:15] if line.strip()]

    # Reject non-ORD indicator lines: tests, questionnaires, exam tickets, labs, syllabus, contracts, agreements, tech specs, resumes, reports
    is_non_ord_context = bool(re.search(
        r"\b(?:тест|вопросы\s+к|контрольные\s+вопросы|экзаменационн|зачет[ае]?|билет\s+№|семестр|лабораторн|практическ|учебн|"
        r"договор|соглашени[ея]|коммерческ[а-я]*\s+предложен|бизнес-план|паспорт\s+проекта|справк[аиеу]|"
        r"техническое\s+задание|\bтз\b|резюме(?:\s+кандидата)?|отчет\s+о|отчет\s+об)\b",
        "\n".join(candidate_paras[:5]),
        re.I,
    ))

    if not is_non_ord_context:
        for p_raw in candidate_paras:
            p_up = p_raw.upper()
            # Skip lines that are questions, list items, or course headers
            if re.match(r"^\d+[\.\)]", p_raw) or any(k in p_up for k in ("ВОПРОС", "ТЕСТ", "БИЛЕТ", "ЗАДАНИЕ", "ТЕМА:", "ПРЕПОДАВАТЕЛЬ")):
                continue

            # Check if this paragraph is a standalone kind or starts with kind + title/number
            if re.match(r"^ПРИКАЗ\b", p_up):
                doc_kind = "Приказ"
                break
            elif re.match(r"^РАСПОРЯЖЕНИЕ\b", p_up):
                doc_kind = "Распоряжение"
                break
            elif re.match(r"^СЛУЖЕБНАЯ\s+ЗАПИСКА\b", p_up) or "СЛУЖЕБНАЯ ЗАПИСКА" in p_up:
                doc_kind = "Служебная записка"
                break
            elif re.match(r"^ПРОТОКОЛ\b", p_up):
                doc_kind = "Протокол"
                break
            elif re.match(r"^АКТ\b", p_up):
                doc_kind = "Акт"
                break
            elif re.match(r"^(?:ДОГОВОР|СОГЛАШЕНИЕ)\b", p_up) or "ДОГОВОР" in p_up or "СОГЛАШЕНИЕ" in p_up:
                doc_kind = "Неопознанный документ / Не является ОРД"
                break
            elif re.match(r"^(?:ТЕХНИЧЕСКОЕ\s+ЗАДАНИЕ|ТЗ)\b", p_up) or "ТЕХНИЧЕСКОЕ ЗАДАНИЕ" in p_up:
                doc_kind = "Неопознанный документ / Не является ОРД"
                break
            elif re.match(r"^РЕЗЮМЕ\b", p_up) or p_up == "РЕЗЮМЕ":
                doc_kind = "Неопознанный документ / Не является ОРД"
                break
            elif re.match(r"^ОТЧЕТ\b", p_up) or p_up == "ОТЧЕТ":
                doc_kind = "Неопознанный документ / Не является ОРД"
                break

    # Secondary check: if header didn't match, check for explicit operative formulas in document body
    if doc_kind == "Неопознанный документ / Не является ОРД" and not is_non_ord_context:
        if re.search(r"\bПРИКАЗЫВАЮ\b", up) and re.search(r"\bПРИКАЗ\b", up):
            doc_kind = "Приказ"
        elif re.search(r"\bОБЯЗЫВАЮ\b", up) and re.search(r"\bРАСПОРЯЖЕНИЕ\b", up):
            doc_kind = "Распоряжение"
        elif re.search(r"\bСЛУЖЕБНАЯ\s+ЗАПИСКА\b", up) and re.search(r"\bПРОШУ\b", up):
            doc_kind = "Служебная записка"
        elif (
            re.search(r"\b(?:СОСТАВИЛИ\s+НАСТОЯЩИЙ\s+АКТ|НАСТОЯЩИЙ\s+АКТ\s+СОСТАВЛЕН|КОМИССИЯ\s+В\s+СОСТАВЕ)\b", up)
            and re.search(r"\bАКТ\b", up)
        ):
            doc_kind = "Акт"

    low = all_text.lower()
    module_code: Optional[str] = None
    if doc_kind != "Неопознанный документ / Не является ОРД":
        if "вводе в эксплуатацию" in low and "информационной системы" in low:
            module_code = "MOD-ORDER-IS"
        elif "назначении ответственных за ведение делопроизводства" in low or "ответственных за делопроизводство" in low:
            module_code = "MOD-ORDER-DOU"
        elif "приеме и передаче дел" in low:
            module_code = "MOD-ORDER-PROP"
        elif "утверждении и введении в действие" in low:
            module_code = "MOD-ORDER-LNA"
        elif "внутреннего аудита" in low:
            module_code = "MOD-DISP-AUDIT"
        elif "проектного комитета" in low:
            module_code = "MOD-DISP-PROJECT"
        elif "рабочей группы" in low:
            module_code = "MOD-DISP-WORKGROUP"

    return doc_kind, module_code


def _is_ord_document(doc: docx.Document, all_text: str, doc_kind: str) -> bool:
    """Validate whether the uploaded document is an official organizational-administrative document (ОРД).

    Checks 4 mandatory core requisites (OD-184 / GOST R 7.0.97-2016):
    1. Document kind in header / title zone (Приказ, Распоряжение, Служебная записка, Акт, etc.).
    2. Organization requisites (Организация, ООО, АО, ПАО, Департамент, Филиал, etc.).
    3. Administrative / dispositive / request formula (ПРИКАЗЫВАЮ, ОБЯЗЫВАЮ, ПРОШУ, В целях..., etc.).
    4. Signature requisite (position + FIO / signature placeholder / visualization EP).
    """
    if doc_kind == "Неопознанный документ / Не является ОРД":
        return False

    up_text = all_text.upper()

    # Pre-filter: educational / test / non-ORD materials are never corporate ORD
    header_sample = "\n".join(p.text.strip() for p in doc.paragraphs[:10] if p.text.strip())
    if re.search(
        r"\b(?:тест|вопросы\s+к|контрольные\s+вопросы|экзаменационн|зачет[ае]?|билет\s+№|семестр|лабораторн|практическ|учебн)\b",
        header_sample,
        re.I,
    ):
        return False

    # 1. Kind in header: validated by doc_kind detection
    has_kind_in_header = doc_kind in ("Приказ", "Распоряжение", "Служебная записка", "Акт", "Протокол")

    # 2. Organization requisites
    has_org = bool(re.search(
        r"\b(?:организаци[яиейю]|обществ[оа]|компани[яиейю]|министерств[оа]|учрежден[ие][яейю]|департамент|управлени[ея]|филиал|ооо|ао|пао|зао|предприят[ие][яейю]|наименование\s+организации)\b",
        all_text,
        re.I,
    )) or ("НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ" in up_text)

    # 3. Administrative / dispositive formula
    has_operative = bool(re.search(
        r"\b(?:приказываю|обязываю|постановляю|постановили|решили|слушали|выступили|утверждаю|утвердить|прошу(?:\s+согласовать|\s+выделить|\s+разрешить|\s+рассмотреть)?|в\s+целях|в\s+соответствии\s+с|во\s+исполнение|в\s+связи\s+с|на\s+основании|составили\s+настоящий\s+акт|составлен\s+акт|комиссия\s+в\s+составе|проведена\s+проверка|контроль\s+за\s+исполнением|настоящий\s+акт)\b",
        all_text,
        re.I,
    ))

    # 4. Signature block (paper signature or electronic visualization stamp)
    has_position = bool(re.search(
        r"(?:директор|начальник|руководитель|заместитель|председатель|управляющий|заведующий|главный|президент|генеральный|составитель|секретарь)",
        all_text,
        re.I,
    ))
    has_fio = bool(re.search(
        r"(?:[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁа-яё]+|[А-ЯЁа-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.|И\.\s*О\.\s*Фамилия|\(подпись\))",
        all_text,
    ))
    has_signature = (has_position and has_fio) or ("ВИЗУАЛИЗАЦИЯ ЭП" in up_text)

    # An official ORD must have:
    # - Valid document kind in header
    # - Substantive dispositive text (has_operative) OR signature block (has_signature)
    # - At least 2 of the 4 requisites in total
    if not has_kind_in_header:
        return False

    if not (has_operative or has_signature):
        return False

    requisites_count = sum([has_kind_in_header, has_org, has_operative, has_signature])
    return requisites_count >= 2


def remediate_docx(doc_bytes: bytes, file_name: str) -> Optional[str]:
    """Auto-Fix routine: cleans template artifacts, standardizes margins, fonts, line spacing, indents.

    Saves cleaned and properly formatted document to generated_docs/fixed_{uuid}.docx
    and returns relative download URL '/v1/documents/download/fixed_{uuid}.docx'.
    """
    try:
        doc = docx.Document(io.BytesIO(doc_bytes))

        # 1. Standardize page margins (Section 5.1 OD-184: left 30mm, right 15mm, top 20mm, bottom 20mm)
        for section in doc.sections:
            section.left_margin = Mm(30)
            section.right_margin = Mm(15)
            section.top_margin = Mm(20)
            section.bottom_margin = Mm(20)

        # 2. Iterate paragraphs to clean text and standardize formatting
        paras_to_remove = []
        in_notes_block = False
        for p_idx, p in enumerate(doc.paragraphs):
            raw_text = p.text
            up = raw_text.upper().strip()

            # Remove technical marker 'МОДУЛЬНЫЙ ДОКУМЕНТ'
            if "МОДУЛЬНЫЙ ДОКУМЕНТ" in up:
                paras_to_remove.append(p)
                continue

            # Remove template approval stamp 'УТВЕРЖДЕНО... УТВ-...'
            if "УТВ-" in up and re.search(r"\bУТВЕРЖД(?:ЕНО|АЮ)\b", up):
                paras_to_remove.append(p)
                continue

            # Check if this paragraph starts the explanatory notes section at the end of the document
            if (
                "ПРИМЕЧАНИЯ ПО ЗАПОЛНЕНИЮ" in up
                or "ПОЯСНЕНИЯ ПО ЗАПОЛНЕНИЮ" in up
                or "ПРИМЕЧАНИЯ К ШАБЛОНУ" in up
                or "ПОЯСНЕНИЯ К ШАБЛОНУ" in up
                or up.startswith("ПРИМЕЧАНИЯ:")
                or up == "ПРИМЕЧАНИЯ"
            ):
                in_notes_block = True

            if in_notes_block:
                paras_to_remove.append(p)
                continue

            # Text cleaning: remove footnote brackets [1]..[99], clean empty placeholders
            cleaned = re.sub(r"\[\d+\]", "", raw_text)
            cleaned = re.sub(r"_{2,}", "[Заполнить: реквизит]", cleaned)
            cleaned = re.sub(r"\[Указать\s+([^\]]+)\]", r"[Заполнить: \1]", cleaned, flags=re.I)

            # Title remediation: strip trailing dot, ensure no underline, no first line indent (p. 5.5 OD-184)
            is_title = bool(p_idx <= 12 and re.match(r"^(?:О|Об)\s+[а-яА-ЯёЁa-zA-Z0-9]", cleaned.strip(), re.I))
            if is_title:
                cleaned = cleaned.rstrip(".")

            if cleaned != raw_text:
                p.text = cleaned

            # Formatting: Times New Roman 13 pt, line spacing 1.15
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)

            # First line indent 1.25 cm for body paragraphs, 0 cm for title
            is_centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
            is_short = len(cleaned.strip()) < 35 or cleaned.strip().isupper()
            if is_title:
                p.paragraph_format.first_line_indent = Cm(0)
            elif not is_centered and not is_short:
                p.paragraph_format.first_line_indent = Cm(1.25)

            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                if is_title:
                    r.font.underline = False
                try:
                    rPr = r._element.get_or_add_rPr()
                    rFonts = rPr.get_or_add_rFonts()
                    rFonts.set(docx.oxml.ns.qn("w:cs"), "Times New Roman")
                    rFonts.set(docx.oxml.ns.qn("w:eastAsia"), "Times New Roman")
                except Exception:
                    pass

        # Remove flagged paragraphs from document XML
        for p in paras_to_remove:
            try:
                p._element.getparent().remove(p._element)
            except Exception:
                p.text = ""

        # 3. Clean tables
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        c_text = p.text
                        c_clean = re.sub(r"\[\d+\]", "", c_text)
                        c_clean = re.sub(r"_{2,}", "[Заполнить]", c_clean)
                        if c_clean != c_text:
                            p.text = c_clean
                        p.paragraph_format.line_spacing = 1.0
                        for r in p.runs:
                            r.font.name = "Times New Roman"
                            r.font.size = Pt(11)
                            try:
                                rPr = r._element.get_or_add_rPr()
                                rFonts = rPr.get_or_add_rFonts()
                                rFonts.set(docx.oxml.ns.qn("w:cs"), "Times New Roman")
                            except Exception:
                                pass

        # 4. Save to generated_docs/fixed_{uuid}.docx
        settings.GENERATED_DOCS_DIR.mkdir(parents=True, exist_ok=True)
        fixed_uid = uuid.uuid4().hex[:12]
        fixed_file_name = f"fixed_{fixed_uid}.docx"
        fixed_path = settings.GENERATED_DOCS_DIR / fixed_file_name
        doc.save(str(fixed_path))

        return f"/v1/documents/download/{fixed_file_name}"
    except Exception as e:
        logger.warning(f"Failed to auto-fix docx '{file_name}': {e}")
        return None


async def check_docx_compliance(
    session: AsyncSession,
    file_bytes: bytes,
    file_name: str,
) -> ComplianceCheckResponse:
    """Run comprehensive 16-point normative compliance inspection on an uploaded DOCX draft."""
    doc = docx.Document(io.BytesIO(file_bytes))
    para_lines, table_cells = _extract_docx_lines_and_tables(doc)
    all_text = "\n".join(para_lines + table_cells)
    up_text = all_text.upper()

    doc_kind, module_code = _detect_doc_kind_and_module(all_text, doc=doc)
    is_ord = _is_ord_document(doc, all_text, doc_kind)
    is_modular = bool((module_code or "МОДУЛЬНЫЙ ДОКУМЕНТ" in up_text) and is_ord)
    issues: List[ComplianceIssue] = []

    # =========================================================================
    # Group 1: Physical Formatting Rules (Section 5 OD-184)
    # =========================================================================

    # 1. Page Margins (cl. 5.1 OD-184: left >= 30mm, right >= 15mm, top/bottom >= 20mm)
    for i, section in enumerate(doc.sections, 1):
        left_mm = section.left_margin.mm if section.left_margin is not None else 25.4
        right_mm = section.right_margin.mm if section.right_margin is not None else 25.4
        top_mm = section.top_margin.mm if section.top_margin is not None else 25.4
        bottom_mm = section.bottom_margin.mm if section.bottom_margin is not None else 25.4

        margin_errors = []
        if left_mm < 29.5:
            margin_errors.append(f"левое {left_mm:.1f} мм (норма ≥ 30 мм)")
        if right_mm < 14.5:
            margin_errors.append(f"правое {right_mm:.1f} мм (норма ≥ 15 мм)")
        if top_mm < 19.5:
            margin_errors.append(f"верхнее {top_mm:.1f} мм (норма ≥ 20 мм)")
        if bottom_mm < 19.5:
            margin_errors.append(f"нижнее {bottom_mm:.1f} мм (норма ≥ 20 мм)")

        if margin_errors:
            issues.append(
                ComplianceIssue(
                    rule_id="CHK-MARG",
                    severity="warning",
                    clause_ref="П. 5.1 Инструкции ОД-184",
                    title="Нарушение регламентных полей страницы",
                    description=f"В разделе {i} обнаружены некорректные поля: {', '.join(margin_errors)}.",
                    recommendation="Установите поля: левое — не менее 30 мм, правое — не менее 15 мм, верхнее и нижнее — не менее 20 мм (п. 5.1 ОД-184).",
                )
            )
            break

    # 2. Font Family and Size (cl. 5.2 OD-184: Times New Roman, Arial, PT Astra Serif, 12-14 pt)
    default_font, default_size = _extract_doc_defaults(doc)
    fallback_doc_font = default_font or "Calibri"
    fallback_doc_size = default_size if default_size is not None else 11.0

    bad_fonts = set()
    bad_sizes = set()
    for p in doc.paragraphs:
        txt = p.text.strip()
        if len(txt) < 20:
            continue
        style_font = getattr(p.style.font, "name", None) if p.style else None
        style_size = getattr(p.style.font, "size", None) if p.style else None

        for r in p.runs:
            if not r.text.strip():
                continue
            fn = r.font.name or style_font or fallback_doc_font
            fs = r.font.size or style_size
            pt_val = fs.pt if fs is not None else fallback_doc_size

            if fn and fn.strip().lower() not in ALLOWED_FONTS:
                bad_fonts.add(fn.strip())
            if pt_val < 11.5 or pt_val > 14.5:
                bad_sizes.add(f"{pt_val:.1f} pt")

    if bad_fonts or bad_sizes:
        desc_parts = []
        if bad_fonts:
            desc_parts.append(f"нерегламентный шрифт ({', '.join(sorted(bad_fonts)[:3])})")
        if bad_sizes:
            desc_parts.append(f"нерегламентный кегль ({', '.join(sorted(bad_sizes)[:3])})")
        issues.append(
            ComplianceIssue(
                rule_id="CHK-FONT",
                severity="warning",
                clause_ref="П. 5.2 Инструкции ОД-184",
                title="Не регламентированный шрифт или размер кегля",
                description=f"В основном тексте документа обнаружены {', '.join(desc_parts)}.",
                recommendation="Используйте гарнитуры Times New Roman, Arial или PT Astra Serif размером 12–14 pt (п. 5.2 ОД-184).",
            )
        )

    # 3. Line Spacing (cl. 5.3 OD-184: 1.0 - 1.5)
    bad_spacings = set()
    for p in doc.paragraphs:
        if len(p.text.strip()) < 30:
            continue
        ls = p.paragraph_format.line_spacing
        if ls is not None:
            if isinstance(ls, (int, float)):
                if ls < 0.95 or ls > 1.55:
                    bad_spacings.add(f"{ls:.2f}")
            elif hasattr(ls, "pt"):
                if ls.pt < 11.0 or ls.pt > 22.0:
                    bad_spacings.add(f"{ls.pt:.1f} pt")

    if bad_spacings:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-LINE",
                severity="warning",
                clause_ref="П. 5.3 Инструкции ОД-184",
                title="Не регламентированный межстрочный интервал",
                description=f"В тексте обнаружен межстрочный интервал, выходящий за пределы 1.0–1.5: {', '.join(sorted(bad_spacings)[:3])}.",
                recommendation="Установите одинарный или полуторный межстрочный интервал (1.0–1.5, рекомендуется 1.15) (п. 5.3 ОД-184).",
            )
        )

    # 4. Paragraph Indent (cl. 5.3 OD-184: 1.25 cm first line indent)
    body_paras_without_indent = 0
    total_body_paras = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if len(txt) >= 40 and not txt.startswith(("ПРИКАЗ", "РАСПОРЯЖЕНИЕ", "УТВЕРЖД", "О ", "Об ", "Приложение")):
            if p.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT):
                continue
            total_body_paras += 1
            fli = p.paragraph_format.first_line_indent
            if fli is None or abs(fli.cm - 1.25) > 0.35:
                body_paras_without_indent += 1

    if total_body_paras >= 2 and (body_paras_without_indent / total_body_paras) >= 0.5:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-INDENT",
                severity="warning",
                clause_ref="П. 5.3 Инструкции ОД-184",
                title="Отсутствует абзацный отступ первой строки (красная строка 1.25 см)",
                description="В абзацах основного текста не соблюдается обязательный абзацный отступ 1.25 см.",
                recommendation="Установите отступ первой строки каждого абзаца ровно 1.25 см (без пробелов и табуляции) (п. 5.3 ОД-184).",
            )
        )

    # 5. Document Title Formatting (cl. 5.3, 5.5 OD-184)
    title_paras = []
    for p in doc.paragraphs[:12]:
        t = p.text.strip()
        if re.match(r"^(?:О|Об)\s+[а-яА-ЯёЁa-zA-Z0-9]", t, re.I):
            title_paras.append((p, t))

    if not title_paras:
        if doc_kind in ("Приказ", "Распоряжение", "Служебная записка"):
            issues.append(
                ComplianceIssue(
                    rule_id="CHK-TITLE",
                    severity="warning",
                    clause_ref="П. 5.3, 5.5 Инструкции ОД-184",
                    title="Отсутствует заголовок к тексту документа (тема «О ...» / «Об ...»)",
                    description="Заголовок к тексту должен кратко формулировать тему документа и начинаться с предлога «О» или «Об».",
                    recommendation="Добавьте заголовок к тексту, начинающийся со слова «О ...» или «Об ...» (п. 5.5 ОД-184).",
                )
            )
    else:
        p_title, t_title = title_paras[0]
        title_issues = []
        if t_title.endswith("."):
            title_issues.append("заголовок заканчивается точкой (точка не ставится)")
        has_underline = any(r.font.underline for r in p_title.runs if r.font.underline)
        if has_underline or "___" in t_title:
            title_issues.append("заголовок подчеркнут (подчеркивание не допускается)")
        if title_issues:
            issues.append(
                ComplianceIssue(
                    rule_id="CHK-TITLE",
                    severity="warning",
                    clause_ref="П. 5.5 Инструкции ОД-184",
                    title="Нарушение оформления заголовка к тексту документа",
                    description=f"В заголовке выявлены нарушения: {'; '.join(title_issues)}.",
                    recommendation="Заголовок к тексту не должен содержать точку на конце и не должен подчеркиваться (п. 5.5 ОД-184).",
                )
            )

    # 6. Signature Requisition (cl. 6.1, 6.3 OD-184)
    has_position = bool(re.search(
        r"(?:директор|начальник|руководитель|заместитель|председатель|управляющий|заведующий|главный|президент|генеральный|составитель|секретарь)",
        all_text,
        re.I,
    ))
    has_fio = bool(re.search(
        r"(?:[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁа-яё]+|[А-ЯЁа-яё]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.|И\.\s*О\.\s*Фамилия)",
        all_text,
    ))
    if not (has_position and has_fio):
        issues.append(
            ComplianceIssue(
                rule_id="CHK-SIGN",
                severity="critical",
                clause_ref="П. 6.1, 6.3 Инструкции ОД-184",
                title="Неполный или отсутствующий реквизит «Подпись»",
                description="В документе не обнаружен полный реквизит подписи уполномоченного лица (наименование должности и расшифровка ФИО).",
                recommendation="Оформите реквизит «Подпись» по п. 6.1 ОД-184: наименование должности, личная подпись / отметка об ЭП, расшифровка инициалов и фамилии.",
            )
        )

    # =========================================================================
    # Group 2: Modular Document & Template Cleanup Rules (Section 3.3)
    # =========================================================================

    # Check 1: Leftover 'МОДУЛЬНЫЙ ДОКУМЕНТ' marker (Section 3.3 Methodological Guidelines)
    if "МОДУЛЬНЫЙ ДОКУМЕНТ" in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-01",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций",
                title="Не удалена техническая отметка «МОДУЛЬНЫЙ ДОКУМЕНТ»",
                description="В верхней части проекта осталась служебная отметка шаблона «МОДУЛЬНЫЙ ДОКУМЕНТ».",
                recommendation="Удалите строку «МОДУЛЬНЫЙ ДОКУМЕНТ» перед направлением проекта на согласование.",
            )
        )

    # Check 2: Leftover upper approval stamp 'УТВЕРЖДЕНО' / 'УТВЕРЖДАЮ' (Section 3.3)
    if re.search(r"\bУТВЕРЖД(?:ЕНО|АЮ)\b", up_text) and "УТВ-" in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-02",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций",
                title="Не удален верхний гриф утверждения модульного шаблона",
                description="В проекте сохранился гриф утверждения самого модульного шаблона (УТВЕРЖДЕНО / № УТВ-...).",
                recommendation="Полностью удалите верхний гриф утверждения шаблона (п. 3.3 Методических рекомендаций).",
            )
        )

    # Check 3: Unfilled blank lines '____' or bracketed placeholders '[Указать ...]' (Section 3.3)
    if re.search(r"_{2,}|\[[^\d\]][^\]]{2,}\]", all_text):
        issues.append(
            ComplianceIssue(
                rule_id="CHK-03",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций, п. 5.1 Инструкции ОД-184",
                title="Обнаружены незаполненные графы или шаблонные заглушки («____» / «[...]»)",
                description="В тексте документа остались незаполненные прочерки «____» или шаблонные подсказки в скобках из шаблона.",
                recommendation="Заполните все переменные поля конкретными реквизитами и удалите линии подчеркивания и подсказки в скобках.",
            )
        )

    # Check 4: Leftover footnote indices [1], [2], [15] (Section 3.3)
    footnote_hits = re.findall(r"\[\d+\]", all_text)
    if footnote_hits:
        sample = ", ".join(sorted(set(footnote_hits))[:5])
        issues.append(
            ComplianceIssue(
                rule_id="CHK-04",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций",
                title=f"Не удалены цифровые сноски шаблона ({sample})",
                description="В тексте проекта остались ссылки на примечания по заполнению вида [1], [2].",
                recommendation="Удалите все квадратные скобки с номерами сносок из заголовка, текста и таблиц.",
            )
        )

    # Check 5: Leftover explanatory notes block at end of document (Section 3.3)
    if "ПРИМЕЧАНИЯ ПО ЗАПОЛНЕНИЮ" in up_text or "ПОЯСНЕНИЯ ПО ЗАПОЛНЕНИЮ" in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-05",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций",
                title="Не удален справочный раздел «Примечания / Пояснения по заполнению»",
                description="В конце документа остался служебный раздел с инструкциями по заполнению модуля.",
                recommendation="Удалите раздел «Примечания по заполнению» и все пункты пояснений в конце файла.",
            )
        )

    # Check 6: Service zone 'ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)' (Section 3.3)
    if is_modular and doc_kind in ("Приказ", "Распоряжение") and "ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)" not in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-06",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций, п. 5.3 Инструкции ОД-184",
                title="Отсутствует служебная зона «ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)»",
                description="Служебная рамка для автоматической простановки даты и номера в ИИС ДО была удалена или повреждена.",
                recommendation="Восстановите блок «ДАТА / РЕГИСТРАЦИОННЫЙ НОМЕР (НЕ УДАЛЯТЬ)» в верхней таблице документа.",
            )
        )

    # Check 7: Service zone 'ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)' (Section 3.3 & cl. 6.1)
    if is_modular and doc_kind in ("Приказ", "Распоряжение") and "ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)" not in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-07",
                severity="critical",
                clause_ref="Разд. 3.3 Методических рекомендаций, п. 6.1 Инструкции ОД-184",
                title="Отсутствует служебная зона «ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)»",
                description="Служебная зона для штампа электронной подписи отсутствует в блоке подписи.",
                recommendation="Сохраните маркер «ВИЗУАЛИЗАЦИЯ ЭП (НЕ УДАЛЯТЬ)» перед ФИО подписанта.",
            )
        )

    # Check 8: Unfilled signatory placeholder 'И. О. Фамилия'
    if re.search(r"И\.\s*О\.\s*Фамилия", all_text, re.I):
        issues.append(
            ComplianceIssue(
                rule_id="CHK-08",
                severity="critical",
                clause_ref="П. 6.1, 6.3 Инструкции ОД-184",
                title="Не указаны инициалы и фамилия подписанта («И. О. Фамилия»)",
                description="В реквизите «Подпись» осталась шаблонная заглушка «И. О. Фамилия».",
                recommendation="Укажите фактические инициалы и фамилию уполномоченного должностного лица.",
            )
        )

    # Check 9: Control clause presence in orders and dispositions (cl. 9.2.6 / 9.7.6)
    if doc_kind in ("Приказ", "Распоряжение") and "КОНТРОЛЬ ЗА ИСПОЛНЕНИЕМ" not in up_text:
        issues.append(
            ComplianceIssue(
                rule_id="CHK-09",
                severity="warning",
                clause_ref="П. 9.2.6, 9.7.6 Инструкции ОД-184",
                title="Отсутствует пункт о контроле за исполнением документа",
                description="Заключительный пункт распорядительной части не содержит указания на должностное лицо, контролирующее исполнение.",
                recommendation="Добавьте заключительный пункт: «Контроль за исполнением настоящего приказа/распоряжения возложить на ... (или оставляю за собой)».",
            )
        )

    # Check 10: Preamble check (cl. 9.1.4)
    if doc_kind in ("Приказ", "Распоряжение") and not any(
        w in up_text for w in ("В ЦЕЛЯХ", "В СООТВЕТСТВИИ С", "ВО ИСПОЛНЕНИЕ", "В СВЯЗИ С")
    ):
        issues.append(
            ComplianceIssue(
                rule_id="CHK-10",
                severity="warning",
                clause_ref="П. 9.1.4, 9.2.3 Инструкции ОД-184",
                title="Не обнаружена преамбула (обоснование издания акта)",
                description="Констатирующая часть (преамбула) должна начинаться словами «В целях...», «В соответствии с...» или «Во исполнение...».",
                recommendation="Добавьте преамбулу с указанием целей и правовых оснований перед словом ПРИКАЗЫВАЮ / ОБЯЗЫВАЮ.",
            )
        )

    if not is_ord:
        doc_kind = "Неопознанный документ / Не является ОРД"
        module_code = None
        is_modular = False
        issues = [
            ComplianceIssue(
                rule_id="CHK-ORD",
                severity="critical",
                clause_ref="Инструкция по ДОУ № ОД-184, ГОСТ Р 7.0.97-2016",
                title="Документ не идентифицирован как ОРД Организации",
                description=(
                    "Загруженный файл не содержит обязательных реквизитов официального организационно-распорядительного документа "
                    "(отсутствует регламентный вид документа в заголовке, реквизиты организации, распорядительная/констатирующая часть "
                    "или реквизит подписи). Документ не пригоден к согласованию в СЭД/ИИС ДО."
                ),
                recommendation=(
                    "Оформите проект документа на официальном бланке Организации в соответствии с требованиями "
                    "Инструкции по ДОУ № ОД-184 (укажите вид документа, реквизиты автора, констатирующую часть и подпись)."
                ),
            )
        ]

    total_checks = 16
    critical_count = sum(1 for i in issues if i.severity == "critical")
    warning_count = sum(1 for i in issues if i.severity == "warning")
    info_count = sum(1 for i in issues if i.severity == "info")

    if not is_ord:
        # Non-ORD document: score capped at 12.5% (2 passed checks out of 16), no auto-fix, no approval route
        passed_checks = 2
        score_pct = round((passed_checks / total_checks) * 100.0, 1)
        passed = False
        is_compliant = False
        fixed_file_url = None
        approval_rec = None
    else:
        passed_checks = max(0, total_checks - len(issues))
        score_pct = round((passed_checks / total_checks) * 100.0, 1)
        passed = (critical_count == 0)
        is_compliant = (len(issues) == 0)

        # Generate Auto-Fix document
        fixed_file_url = remediate_docx(file_bytes, file_name)

        # Calculate recommended approval route
        route_target = module_code or doc_kind
        calc_res = await calculate_approval_route(
            session,
            ApprovalCalculateRequest(
                doc_type=route_target,
                uses_module=is_modular,
            ),
        )
        approval_rec = ApprovalRecommendation(
            rule_number=calc_res.rule_number,
            doc_type_name=calc_res.doc_type_name,
            stages=calc_res.stages,
            signatory=calc_res.signatory,
        )

    return ComplianceCheckResponse(
        file_name=file_name,
        detected_doc_kind=doc_kind,
        is_modular=is_modular,
        detected_module_code=module_code,
        score_pct=score_pct,
        passed=passed,
        is_compliant=is_compliant,
        total_checks=total_checks,
        passed_checks=passed_checks,
        critical_count=critical_count,
        warning_count=warning_count,
        info_count=info_count,
        issues=issues,
        approval_recommendation=approval_rec,
        fixed_file_url=fixed_file_url,
    )
