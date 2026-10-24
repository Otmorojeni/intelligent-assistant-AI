"""Hot-reload document versioning and lifecycle management.

Handles dynamic upload of new regulatory document editions (PDF, DOCX),
automatic archival of previous editions (is_active = False), immediate live
tsvector indexing in PostgreSQL without server restarts, active version rollback,
and document version listing.
"""

from datetime import datetime, timezone
import hashlib
import logging
from pathlib import Path
import re
from typing import Any, Dict, List, Optional
import uuid

import docx
from fastapi import HTTPException, status
import pymupdf
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.config import settings
from src.core.extractor import (
    ExtractedChunk,
    ExtractedDocument,
    clean_pdf_text_line,
    extract_guidelines_document,
)
from src.db.models import AuditLog, DocumentChunk, KnowledgeDocument
from src.schemas.api_models import AuditLogItem, DocumentVersionItem

logger = logging.getLogger(__name__)

UPLOADS_DIR = settings.GENERATED_DOCS_DIR / "uploads"
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)


def compute_file_hash(path: Path) -> str:
    """Calculate SHA-256 checksum of local file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def parse_uploaded_document(
    file_path: Path,
    doc_code: str,
    title: str,
    version: str,
    doc_type: str = "instruction",
    access_level: str = "public",
) -> ExtractedDocument:
    """Parse uploaded PDF or DOCX file on the fly into structured ExtractedDocument and chunks."""
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix not in [".pdf", ".docx"]:
        raise ValueError(
            f"Unsupported file extension '{suffix}'. Only .pdf and .docx are supported."
        )

    file_size = file_path.stat().st_size
    sha256 = compute_file_hash(file_path)
    chunks: List[ExtractedChunk] = []

    if suffix == ".pdf":
        try:
            doc = pymupdf.open(file_path)
        except Exception as e:
            raise ValueError(f"Failed to open PDF document: {e}") from e

        if doc.is_encrypted:
            doc.close()
            raise ValueError("Uploaded PDF is password protected or encrypted.")

        with doc:
            current_chapter = "Титульный лист и общие положения"
            current_clause = "1.1"
            current_page = 1
            current_lines: List[str] = []

            for p_idx in range(len(doc)):
                page_num = p_idx + 1
                page_text = doc[p_idx].get_text().replace("\xa0", " ").replace("\xad", "-")
                lines = [clean_pdf_text_line(l) for l in page_text.split("\n")]
                lines = [l for l in lines if l]

                for line in lines:
                    m_chap = re.match(r"^(?:Глава|Раздел|Статья|[1-9]\d*\.)\s*(\d+)?\.?\s*(.*)", line, re.I)
                    m_cl = re.match(r"^(?:(?:п\.|пп\.|пункт)\s*)?(\d+\.\d+(?:\.\d+)*)\.?\s*(.*)", line, re.I)

                    if m_cl:
                        if current_lines:
                            chunks.append(
                                ExtractedChunk(
                                    chapter=current_chapter,
                                    clause=current_clause,
                                    page_number=current_page,
                                    content="\n".join(current_lines),
                                    metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                                )
                            )
                            current_lines = []
                        current_clause = m_cl.group(1)
                        current_page = page_num
                        current_lines.append(line)
                    elif m_chap:
                        if current_lines:
                            chunks.append(
                                ExtractedChunk(
                                    chapter=current_chapter,
                                    clause=current_clause,
                                    page_number=current_page,
                                    content="\n".join(current_lines),
                                    metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                                )
                            )
                            current_lines = []
                        current_chapter = line
                        ch_num = m_chap.group(1) if m_chap.group(1) else "1"
                        current_clause = f"{ch_num}.0"
                        current_page = page_num
                    else:
                        current_lines.append(line)

            if current_lines:
                chunks.append(
                    ExtractedChunk(
                        chapter=current_chapter,
                        clause=current_clause,
                        page_number=current_page,
                        content="\n".join(current_lines),
                        metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                    )
                )

    elif suffix == ".docx":
        if doc_type in ["guideline", "guidelines"]:
            try:
                extracted = extract_guidelines_document(file_path)
                extracted.doc_code = doc_code
                extracted.title = title
                extracted.version = version
                extracted.doc_type = doc_type
                extracted.access_level = access_level
                for c in extracted.chunks:
                    c.metadata["doc_code"] = doc_code
                    c.metadata["version"] = version
                return extracted
            except Exception as e:
                logger.warning(f"Guidelines extractor failed, falling back to generic docx parser: {e}")

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise ValueError(f"Failed to open DOCX document: {e}") from e

        current_chapter = "Общие положения"
        current_clause = "1.1"
        current_lines: List[str] = []

        for p in doc.paragraphs:
            text = p.text.replace("\xa0", " ").strip()
            if not text:
                continue

            m_chap = re.match(r"^(?:Глава|Раздел|Статья|[1-9]\d*\.)\s*(\d+)?\.?\s*(.*)", text, re.I)
            m_cl = re.match(r"^(?:(?:п\.|пп\.|пункт)\s*)?(\d+\.\d+(?:\.\d+)*)\.?\s*(.*)", text, re.I)

            if m_cl:
                if current_lines:
                    chunks.append(
                        ExtractedChunk(
                            chapter=current_chapter,
                            clause=current_clause,
                            page_number=1,
                            content="\n".join(current_lines),
                            metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                        )
                    )
                    current_lines = []
                current_clause = m_cl.group(1)
                current_lines.append(text)
            elif m_chap and not m_cl:
                if current_lines:
                    chunks.append(
                        ExtractedChunk(
                            chapter=current_chapter,
                            clause=current_clause,
                            page_number=1,
                            content="\n".join(current_lines),
                            metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                        )
                    )
                    current_lines = []
                current_chapter = text
                ch_num = m_chap.group(1) if m_chap.group(1) else "1"
                current_clause = f"{ch_num}.0"
            else:
                current_lines.append(text)

        # Also extract text from all DOCX tables so tabular matrices and forms are indexed
        for tbl_idx, tbl in enumerate(doc.tables, start=1):
            for row in tbl.rows:
                row_cells: List[str] = []
                for cell in row.cells:
                    cell_txt = cell.text.replace("\xa0", " ").strip()
                    if cell_txt and (not row_cells or row_cells[-1] != cell_txt):
                        row_cells.append(cell_txt)
                if row_cells:
                    current_lines.append(" | ".join(row_cells))

        if current_lines:
            chunks.append(
                ExtractedChunk(
                    chapter=current_chapter,
                    clause=current_clause,
                    page_number=1,
                    content="\n".join(current_lines),
                    metadata={"doc_code": doc_code, "version": version, "clause": current_clause},
                )
            )

    has_text = any(bool(c.content.strip()) for c in chunks)
    if not has_text:
        raise ValueError("Uploaded document contains no extractable text or content.")

    return ExtractedDocument(
        doc_code=doc_code,
        title=title,
        doc_type=doc_type,
        version=version,
        access_level=access_level,
        file_name=file_path.name,
        file_size=file_size,
        sha256=sha256,
        chunks=chunks,
    )


async def upload_document_version(
    session: AsyncSession,
    file_bytes: bytes,
    file_name: str,
    doc_code: str,
    title: str,
    version: str,
    doc_type: str = "instruction",
    access_level: str = "public",
) -> KnowledgeDocument:
    """Save uploaded document edition, archive previous versions, parse chunks, and commit hot-reload."""
    from src.core.indexer import save_extracted_document

    clean_code = doc_code.strip()
    clean_version = version.strip()
    clean_title = title.strip()
    clean_type = doc_type.strip()
    clean_access = access_level.strip().lower()
    if clean_access in ("internal", "дсп", "для служебного пользования"):
        clean_access = "dsp"

    safe_filename = Path(file_name).name
    target_path = UPLOADS_DIR / f"{clean_code}_v{clean_version}_{safe_filename}"
    with open(target_path, "wb") as f:
        f.write(file_bytes)

    # 1. Parse the uploaded file on the fly FIRST before making any DB changes!
    try:
        extracted_doc = parse_uploaded_document(
            file_path=target_path,
            doc_code=clean_code,
            title=clean_title,
            version=clean_version,
            doc_type=clean_type,
            access_level=clean_access,
        )
    except Exception as e:
        if target_path.exists():
            try:
                target_path.unlink()
            except Exception:
                pass
        raise ValueError(f"Failed to parse uploaded document: {e}") from e

    # 2. Lock and fetch existing documents with the same doc_code
    stmt = (
        select(KnowledgeDocument)
        .where(KnowledgeDocument.doc_code == clean_code)
        .with_for_update()
    )
    existing_docs = (await session.execute(stmt)).scalars().all()

    # 3. Archive previous versions or remove duplicate same version
    for prev_doc in existing_docs:
        if prev_doc.version == clean_version:
            # Overwriting the exact same version edition
            await session.delete(prev_doc)
        else:
            prev_doc.is_active = False

    await session.flush()

    # 4. Save new document and chunks with search_vector computed column
    db_doc = await save_extracted_document(session, extracted_doc)

    # 5. If an updated Approval Matrix DOCX was uploaded, dynamically refresh new approval_rules
    if (clean_code == "APPROVAL-MATRIX" or clean_type == "matrix") and target_path.suffix.lower() == ".docx":
        try:
            async with session.begin_nested():
                from src.core.extractor import extract_approval_matrix_data
                from src.core.indexer import save_approval_rules
                from src.db.models import ApprovalRule

                _, parsed_rules = extract_approval_matrix_data(target_path)
                existing_nums = set(
                    (await session.execute(select(ApprovalRule.rule_number))).scalars().all()
                )
                new_rules = [r for r in parsed_rules if r.rule_number not in existing_nums]
                if new_rules:
                    await save_approval_rules(session, new_rules)
        except Exception as exc:
            logger.warning(f"Optional matrix rules refresh skipped: {exc}")

    await session.commit()
    await session.refresh(db_doc)

    logger.info(
        f"Uploaded and activated new document edition: {clean_code} v{clean_version} "
        f"({len(extracted_doc.chunks)} chunks, access={clean_access})"
    )
    return db_doc


async def activate_document_version(
    session: AsyncSession,
    doc_code: str,
    version: str,
) -> KnowledgeDocument:
    """Switch active document edition: sets target to is_active=True, others to is_active=False."""
    clean_code = (doc_code or "").strip()
    clean_version = (version or "").strip()

    stmt = (
        select(KnowledgeDocument)
        .where(KnowledgeDocument.doc_code == clean_code)
        .with_for_update()
    )
    docs = (await session.execute(stmt)).scalars().all()

    if not docs:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Document with code '{clean_code}' not found.",
        )

    target_doc: Optional[KnowledgeDocument] = None
    for d in docs:
        if d.version == clean_version:
            target_doc = d
            d.is_active = True
        else:
            d.is_active = False

    if not target_doc:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Version '{clean_version}' for document '{clean_code}' not found.",
        )

    await session.commit()
    await session.refresh(target_doc)
    logger.info(f"Activated document edition {clean_code} v{clean_version}")
    return target_doc


async def delete_document_version(
    session: AsyncSession,
    doc_code: str,
    version: Optional[str] = None,
) -> Dict[str, Any]:
    """Delete a document version (or all versions of doc_code if version is None).

    If an active version is deleted and other versions remain, automatically activates
    the latest remaining version.
    """
    clean_code = (doc_code or "").strip()
    if not clean_code:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Document code is required.",
        )

    stmt = select(KnowledgeDocument).where(KnowledgeDocument.doc_code == clean_code)
    if version:
        clean_version = version.strip()
        stmt = stmt.where(KnowledgeDocument.version == clean_version)
    else:
        clean_version = None

    docs_to_delete = (await session.execute(stmt)).scalars().all()
    if not docs_to_delete:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Document '{clean_code}'" + (f" version '{clean_version}'" if clean_version else "") + " not found.",
        )

    deleted_versions = [d.version for d in docs_to_delete]
    was_active = any(d.is_active for d in docs_to_delete)

    for d in docs_to_delete:
        await session.delete(d)

    await session.flush()

    # If an active version was deleted, see if any remaining version exists to activate
    activated_remaining = None
    if was_active and clean_version:
        remaining_stmt = (
            select(KnowledgeDocument)
            .where(KnowledgeDocument.doc_code == clean_code)
            .order_by(KnowledgeDocument.created_at.desc())
        )
        remaining_docs = (await session.execute(remaining_stmt)).scalars().all()
        if remaining_docs and not any(r.is_active for r in remaining_docs):
            remaining_docs[0].is_active = True
            activated_remaining = remaining_docs[0].version

    await session.commit()
    logger.info(
        f"Deleted document {clean_code} versions {deleted_versions}. "
        f"Auto-activated: {activated_remaining}"
    )

    return {
        "deleted_doc_code": clean_code,
        "deleted_versions": deleted_versions,
        "deleted_count": len(deleted_versions),
        "activated_version": activated_remaining,
    }


async def get_all_document_versions(
    session: AsyncSession,
) -> List[DocumentVersionItem]:
    """Retrieve full catalog of all document versions with chunk counts and active statuses."""
    stmt = (
        select(
            KnowledgeDocument,
            func.count(DocumentChunk.id).label("chunk_count"),
        )
        .outerjoin(DocumentChunk, KnowledgeDocument.id == DocumentChunk.document_id)
        .group_by(KnowledgeDocument.id)
        .order_by(KnowledgeDocument.doc_code, KnowledgeDocument.created_at.desc())
    )
    rows = (await session.execute(stmt)).all()

    items: List[DocumentVersionItem] = []
    for doc, count in rows:
        items.append(
            DocumentVersionItem(
                id=str(doc.id),
                doc_code=doc.doc_code,
                title=doc.title,
                doc_type=doc.doc_type,
                version=doc.version,
                access_level=doc.access_level,
                is_active=doc.is_active,
                file_name=doc.file_name,
                file_size=doc.file_size,
                sha256=doc.sha256,
                chunk_count=int(count or 0),
                created_at=doc.created_at,
                updated_at=doc.updated_at,
            )
        )
    return items


async def query_audit_logs(
    session: AsyncSession,
    action_type: Optional[str] = None,
    user_role: Optional[str] = None,
    status_filter: Optional[str] = None,
    limit: int = 50,
    offset: int = 0,
) -> List[AuditLogItem]:
    """Retrieve filtered audit trail log entries from PostgreSQL audit_logs table."""
    stmt = select(AuditLog)
    if action_type:
        stmt = stmt.where(AuditLog.action_type == action_type)
    if user_role:
        stmt = stmt.where(AuditLog.user_role == user_role)
    if status_filter:
        stmt = stmt.where(AuditLog.status == status_filter)

    stmt = stmt.order_by(AuditLog.timestamp.desc()).limit(limit).offset(offset)
    res = await session.execute(stmt)
    logs = res.scalars().all()

    return [
        AuditLogItem(
            id=str(log.id),
            timestamp=log.timestamp,
            user_role=log.user_role,
            user_name=log.user_name,
            action_type=log.action_type,
            query_text=log.query_text,
            response_summary=log.response_summary,
            status=log.status,
            latency_ms=log.latency_ms,
            ip_address=log.ip_address,
        )
        for log in logs
    ]
